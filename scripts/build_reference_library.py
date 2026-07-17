"""
build_reference_library.py — Build the retrieve-and-verify reference library.

Reads:
  - perinatal_knowledge_base/dept_mch_feed/perinatal_misinformation.docx
  - perinatal_knowledge_base/dept_mch_feed/perinatal_comprehensive.docx
  - ml/data/claim_type_map.json        (claim id -> JGIM type)
  - ml/data/claim_paraphrases.json     (claim id -> assertion paraphrases)

Writes:
  - maven_app/anchors/reference_library.json

Entry schema:
  {id, text, kind: misinfo|authority, domain, type_id, correction, parent_id}

Idempotent. Re-run when any input changes. Replaces scripts/build_anchors.py.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path
from typing import Optional

import nltk
from docx import Document
from docx.text.paragraph import Paragraph
from nltk.tokenize import sent_tokenize

nltk.download('punkt', quiet=True)
nltk.download('punkt_tab', quiet=True)

ROOT = Path(__file__).resolve().parents[1]
FEED_DIR = ROOT / 'perinatal_knowledge_base' / 'dept_mch_feed'
MISINFO_DOCX = FEED_DIR / 'perinatal_misinformation.docx'
COMPRE_DOCX = FEED_DIR / 'perinatal_comprehensive.docx'
TYPE_MAP_PATH = ROOT / 'ml' / 'data' / 'claim_type_map.json'
PARAPHRASE_PATH = ROOT / 'ml' / 'data' / 'claim_paraphrases.json'
OUT_PATH = ROOT / 'maven_app' / 'anchors' / 'reference_library.json'

# --- Ported verbatim from scripts/build_anchors.py: ---------------------
# style_name(), is_heading(), normalize(), parse_misinfo_pairs(), and the
# GUIDELINE_RE / SECTION_LETTER_RE / SECTION_TO_DOMAIN constants.
# Copy those function bodies here unchanged (build_anchors.py is in the
# repo until Task 13). Do NOT port extract_authority_sentences() or
# extract_comprehensive_paragraphs() — replaced below.
# -------------------------------------------------------------------------
GUIDELINE_RE = re.compile(
    r'\b(ACOG|ACNM|AWHONN|AAP|SMFM|SFP|NAF|USPSTF|CDC|AHRQ|HRSA|WHO|'
    r'MedlinePlus|NIH|NICHD|FDA|UNICEF|ACIP|LactMed|Cochrane|IOM|NAS)\b'
)
SECTION_LETTER_RE = re.compile(r'^(\d+)([A-Z])\.')

SECTION_TO_DOMAIN = {
    '1A': 'antenatal',
    '1B': 'intrapartum',
    '1C': 'fourth_trimester',
    '1D': 'family_planning',
    '2A': 'clinical_provider',
    '2B': 'clinical_gaslighting',
    '2C': 'clinical_guidelines',
    '3A': 'social_motherhood',
    '3B': 'social_family',
    '3C': 'social_media',
    '3D': 'social_help_seeking',
    '4A': 'info_structural',
    '4B': 'info_topic_misinfo',
    '4C': 'info_spread_mechanisms',
    '4D': 'info_correction_strategies',
    '4E': 'info_quality_signals',
}


def style_name(p: Paragraph) -> str:
    try:
        return p.style.name if p.style else ''
    except Exception:
        return ''


def is_heading(p: Paragraph, level: Optional[int] = None) -> bool:
    sn = style_name(p)
    if not sn.startswith('Heading'):
        return False
    if level is None:
        return True
    return sn == f'Heading {level}'


def normalize(s: str) -> str:
    return re.sub(r'\s+', ' ', s).strip()


def parse_misinfo_pairs(doc_path: Path) -> list[dict]:
    doc = Document(str(doc_path))
    pairs: list[dict] = []
    current_h1: str = ''
    current_h2: str = ''
    current_section_letter: str = ''
    pending_claim: Optional[str] = None

    for p in doc.paragraphs:
        text = normalize(p.text)
        if not text:
            continue

        if is_heading(p, 1):
            current_h1 = text
            current_h2 = ''
            current_section_letter = ''
            pending_claim = None
            continue
        if is_heading(p, 2):
            current_h2 = text
            m = SECTION_LETTER_RE.match(text)
            current_section_letter = f'{m.group(1)}{m.group(2)}' if m else ''
            pending_claim = None
            continue
        if is_heading(p, 3):
            pending_claim = None
            continue

        # CLAIM line
        if text.upper().startswith('CLAIM:'):
            pending_claim = text[len('CLAIM:'):].strip().lstrip("'\"").rstrip("'\"")
            continue

        # EVIDENCE line — must follow a CLAIM
        if text.upper().startswith('EVIDENCE:') and pending_claim:
            evidence = text[len('EVIDENCE:'):].strip()
            domain = SECTION_TO_DOMAIN.get(
                current_section_letter,
                current_h2.lower().replace(' ', '_')[:40] or 'unknown',
            )
            pairs.append({
                'claim':    pending_claim,
                'evidence': evidence,
                'domain':   domain,
            })
            pending_claim = None

    return pairs
# -------------------------------------------------------------------------

ORG_NAME_LINE = re.compile(r'^[A-Z]{2,}(\s*[—–-]|\s+is\s+the\b)')
JUNK_MARKERS = ('Edition', 'Resource for AI', '|')
MIN_WORDS, MAX_WORDS = 8, 60

# Bibliographic / resource-index lines from the comprehensive doc's
# "authoritative guideline bodies" section (e.g. "Key outputs: Position
# Statements..."; "Primary authorities: ACOG, SMFM..."). These pass the
# word-count and heading filters (valid sentence length, not ALL-CAPS,
# don't match ORG_NAME_LINE since the label isn't an acronym) but they
# list publications/orgs/topics rather than assert a checkable fact,
# which makes them poor NLI-verification premises. Found by eyeballing
# a spot-check sample (~8.6% of extracted sentences matched this shape)
# and tightened here per the builder's own filter-quality guidance.
INDEX_LINE_RE = re.compile(
    r'^(Key (outputs|publications|resources|programmes|guidance|topics|areas'
    r'|perinatal topic pages)|Primary (authority|authorities'
    r'|guideline authorities|evidence base)|Other endorsed resources'
    r'|Guideline bodies integrated throughout|Publication types'
    r'|Relevant to perinatal care for|Perinatal recommendations'
    r'|Developmental milestones)\s*:'
)


def extract_authority_statements(doc_path) -> list[str]:
    """Declarative sentences from the comprehensive docx, junk-filtered."""
    from docx import Document
    doc = Document(str(doc_path))
    sentences: list[str] = []
    seen = set()
    for para in doc.paragraphs:
        if is_heading(para):
            continue
        text = normalize(para.text)
        if not text or any(m in text for m in JUNK_MARKERS):
            continue
        for sent in sent_tokenize(text):
            sent = sent.strip()
            words = sent.split()
            if not (MIN_WORDS <= len(words) <= MAX_WORDS):
                continue
            if not sent.endswith('.'):
                continue
            if ORG_NAME_LINE.match(sent):
                continue
            if INDEX_LINE_RE.match(sent):
                continue
            if sent.upper() == sent:  # ALL-CAPS headings
                continue
            if sent not in seen:
                seen.add(sent)
                sentences.append(sent)
    return sentences


def main() -> int:
    pairs = parse_misinfo_pairs(MISINFO_DOCX)  # [{claim, evidence, domain}, ...]

    type_map_doc = json.loads(TYPE_MAP_PATH.read_text(encoding='utf-8'))
    type_map = {c['id']: c for c in type_map_doc['claims']}
    paraphrases = json.loads(PARAPHRASE_PATH.read_text(encoding='utf-8'))

    entries = []

    # Base misinfo claims
    for i, pair in enumerate(pairs, 1):
        cid = f'mis-{i:03d}'
        mapped = type_map.get(cid)
        if mapped is None:
            sys.exit(f'ERROR: {cid} missing from claim_type_map.json')
        if normalize(mapped['claim']) != normalize(pair['claim']):
            sys.exit(f'ERROR: {cid} claim text drift — regenerate claim_type_map.json\n'
                     f'  map: {mapped["claim"]!r}\n  docx: {pair["claim"]!r}')
        entries.append({'id': cid, 'text': pair['claim'], 'kind': 'misinfo',
                        'domain': pair['domain'], 'type_id': mapped['type_id'],
                        'correction': pair['evidence'], 'parent_id': None})

    # Paraphrase expansions
    by_id = {e['id']: e for e in entries}
    for parent_id, texts in paraphrases.items():
        if parent_id.startswith('_'):
            continue
        parent = by_id.get(parent_id)
        if parent is None:
            sys.exit(f'ERROR: paraphrase parent {parent_id} not found')
        for j, text in enumerate(texts, 1):
            entries.append({'id': f'{parent_id}-p{j}', 'text': text,
                            'kind': 'misinfo', 'domain': parent['domain'],
                            'type_id': parent['type_id'],
                            'correction': parent['correction'],
                            'parent_id': parent_id})

    # Authority: every correction is an authority statement...
    auth_texts, seen = [], set()
    for pair in pairs:
        ev = normalize(pair['evidence'])
        if ev and ev not in seen:
            seen.add(ev)
            auth_texts.append(pair['evidence'])
    # ...plus extracted declarative sentences from the comprehensive docx.
    for sent in extract_authority_statements(COMPRE_DOCX):
        if sent not in seen:
            seen.add(sent)
            auth_texts.append(sent)

    for i, text in enumerate(auth_texts, 1):
        entries.append({'id': f'auth-{i:03d}', 'text': text, 'kind': 'authority',
                        'domain': None, 'type_id': None, 'correction': None,
                        'parent_id': None})

    OUT_PATH.write_text(json.dumps(entries, indent=2, ensure_ascii=False),
                        encoding='utf-8')
    n_mis = sum(1 for e in entries if e['kind'] == 'misinfo')
    n_base = sum(1 for e in entries if e['kind'] == 'misinfo' and e['parent_id'] is None)
    n_auth = len(entries) - n_mis
    print(f'Wrote {OUT_PATH}')
    print(f'  misinfo: {n_mis} ({n_base} base + {n_mis - n_base} paraphrases)')
    print(f'  authority: {n_auth}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
